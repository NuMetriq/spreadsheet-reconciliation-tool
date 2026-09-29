import tempfile
import unittest
import subprocess
import sys
from pathlib import Path
from openpyxl import Workbook

from src.reconcile import (
    compare_rows,
    index_rows,
    parse_decimal,
    read_csv,
    read_xlsx,
    read_table,
    reconcile,
    summarize_results,
    write_report,
)


class TestCompareRows(unittest.TestCase):
    def test_identical_rows_have_no_differences(self):
        source = {"customer": "Acme Supply", "amount": "150.00"}
        target = {"customer": "Acme Supply", "amount": "150.00"}

        result = compare_rows(source, target, ["customer", "amount"])

        self.assertEqual(result, {})

    def test_changed_amount_returns_both_values(self):
        source = {"customer": "Acme Supply", "amount": "150.00"}
        target = {"customer": "Acme Supply", "amount": "175.00"}

        result = compare_rows(source, target, ["customer", "amount"])

        self.assertEqual(result, {"amount": ("150.00", "175.00")})

    def test_unselected_columns_are_ignored(self):
        source = {"customer": "Acme Supply", "amount": "150.00"}
        target = {"customer": "Acme Supply", "amount": "175.00"}

        result = compare_rows(source, target, ["customer"])

        self.assertEqual(result, {})

    def test_numeric_comparison_ignores_decimal_format(self):
        source = {"amount": "150.00"}
        target = {"amount": "150.0"}

        result = compare_rows(
            source,
            target,
            ["amount"],
            numeric_columns=["amount"],
        )

        self.assertEqual(result, {})

    def test_text_comparison_preserves_format_differences(self):
        source = {"amount": "150.00"}
        target = {"amount": "150.0"}

        result = compare_rows(source, target, ["amount"])

        self.assertEqual(
            result,
            {"amount": ("150.00", "150.0")},
        )

    def test_numeric_difference_preserves_original_text(self):
        source = {"amount": "150.00"}
        target = {"amount": "175.0"}

        result = compare_rows(
            source,
            target,
            ["amount"],
            numeric_columns=["amount"],
        )

        self.assertEqual(
            result,
            {"amount": ("150.00", "175.0")},
        )

class TestReconcile(unittest.TestCase):
    def test_matches_by_id_when_rows_are_reordered(self):
        source = [
            {"id": "A", "amount": "10.00"},
            {"id": "B", "amount": "20.00"},
        ]
        target = [
            {"id": "B", "amount": "25.00"},
            {"id": "A", "amount": "10.00"},
        ]

        result = reconcile(source, target, "id", ["amount"])

        self.assertEqual(result, [
            {
                "record_id": "A",
                "status": "matched",
                "column": "",
                "source_value": "",
                "target_value": "",
            },
            {
                "record_id": "B",
                "status": "changed",
                "column": "amount",
                "source_value": "20.00",
                "target_value": "25.00",
            },
        ])

    def test_reports_records_missing_from_either_file(self):
        source = [{"id": "A", "amount": "10.00"}]
        target = [{"id": "B", "amount": "20.00"}]

        result = reconcile(source, target, "id", ["amount"])

        self.assertEqual(result, [
            {
                "record_id": "A",
                "status": "source_only",
                "column": "",
                "source_value": "",
                "target_value": "",
            },
            {
                "record_id": "B",
                "status": "target_only",
                "column": "",
                "source_value": "",
                "target_value": "",
            },
        ])

    def test_numeric_comparison_matches_equivalent_amounts(self):
        source = [{"id": "A", "amount": "150.00"}]
        target = [{"id": "A", "amount": "150.0"}]

        result = reconcile(
            source,
            target,
            key_column="id",
            comparison_columns=["amount"],
            numeric_columns=["amount"],
        )

        self.assertEqual(result, [
            {
                "record_id": "A",
                "status": "matched",
                "column": "",
                "source_value": "",
                "target_value": "",
            },
        ])

    def test_rejects_numeric_column_outside_comparison_columns(self):
        source = [{"id": "A", "customer": "Acme", "amount": "150.00"}]
        target = [{"id": "A", "customer": "Acme", "amount": "150.00"}]

        with self.assertRaisesRegex(
            ValueError,
            "Numeric columns must also be comparison columns: amount",
        ):
            reconcile(
                source,
                target,
                key_column="id",
                comparison_columns=["customer"],
                numeric_columns=["amount"],
            )


class TestIndexRows(unittest.TestCase):
    def test_rejects_missing_key_column(self):
        rows = [{"amount": "10.00"}]

        with self.assertRaisesRegex(ValueError, "missing key column"):
            index_rows(rows, "id")

    def test_rejects_blank_key(self):
        rows = [{"id": "   ", "amount": "10.00"}]

        with self.assertRaisesRegex(ValueError, "blank key"):
            index_rows(rows, "id")

    def test_rejects_duplicate_keys(self):
        rows = [
            {"id": "A", "amount": "10.00"},
            {"id": "A", "amount": "20.00"},
        ]

        with self.assertRaisesRegex(ValueError, "Duplicate key"):
            index_rows(rows, "id")


class TestReadCsv(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.csv_path = Path(self.temp_dir.name) / "input.csv"

    def test_reads_valid_csv(self):
        self.csv_path.write_text(
            "id,amount\nA,10.00\n",
            encoding="utf-8",
        )

        result = read_csv(self.csv_path, ["id", "amount"])

        self.assertEqual(result, [{"id": "A", "amount": "10.00"}])

    def test_rejects_missing_column_with_no_data_rows(self):
        self.csv_path.write_text(
            "id,customer\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(
            ValueError, "missing required columns: amount"
        ):
            read_csv(self.csv_path, ["id", "amount"])

    def test_rejects_duplicate_headers(self):
        self.csv_path.write_text(
            "id,amount,amount\nA,10.00,20.00\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "duplicate column names"):
            read_csv(self.csv_path)

    def test_rejects_empty_file(self):
        self.csv_path.write_text("", encoding="utf-8")

        with self.assertRaisesRegex(ValueError, "missing header row"):
            read_csv(self.csv_path)

    def test_rejects_extra_fields(self):
        self.csv_path.write_text(
            "id,amount\nA,10.00,unexpected\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "data row 1 has extra fields"):
            read_csv(self.csv_path)

    def test_rejects_missing_fields(self):
        self.csv_path.write_text(
            "id,amount\nA\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(ValueError, "data row 1 has missing fields"):
            read_csv(self.csv_path)


class TestParseDecimal(unittest.TestCase):
    def test_equivalent_decimal_formats_are_equal(self):
        self.assertEqual(
            parse_decimal("150.00"),
            parse_decimal("150.0"),
        )

    def test_distinguishes_large_values_one_cent_apart(self):
        self.assertNotEqual(
            parse_decimal("1000000000000000.01"),
            parse_decimal("1000000000000000.02"),
        )

    def test_rejects_invalid_numeric_text(self):
        for value in ("", "abc", "$150.00", "1,000.00"):
            with self.subTest(value=value):
                with self.assertRaisesRegex(
                    ValueError, "Invalid numeric value"
                ):
                    parse_decimal(value)

    def test_rejects_nonfinite_values(self):
        for value in ("NaN", "Infinity", "-Infinity"):
            with self.subTest(value=value):
                with self.assertRaisesRegex(
                    ValueError, "must be finite"
                ):
                    parse_decimal(value)


class TestSummarizeResults(unittest.TestCase):
    def test_counts_changed_record_once_for_multiple_differences(self):
        source = [
            {"id": "A", "customer": "Acme", "amount": "10.00"},
        ]
        target = [
            {"id": "A", "customer": "Acme Supply", "amount": "15.00"},
        ]

        results = reconcile(
            source,
            target,
            key_column="id",
            comparison_columns=["customer", "amount"],
        )

        self.assertEqual(len(results), 2)

        summary = summarize_results(results)

        self.assertEqual(summary, {
            "matched": 0,
            "changed": 1,
            "source_only": 0,
            "target_only": 0,
        })

    def test_empty_results_have_zero_counts(self):
        summary = summarize_results([])

        self.assertEqual(summary, {
            "matched": 0,
            "changed": 0,
            "source_only": 0,
            "target_only": 0,
        })


class TestWriteReport(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.output_path = (
            Path(self.temp_dir.name) / "reports" / "result.csv"
        )

    def test_preserves_special_characters_in_report(self):
        results = [
            {
                "record_id": "A",
                "status": "changed",
                "column": "customer",
                "source_value": 'Acme, "North"',
                "target_value": "Café Supply",
            },
        ]

        write_report(results, self.output_path)
        loaded_rows = read_csv(self.output_path)

        self.assertEqual(loaded_rows, results)

    def test_empty_report_still_has_headers(self):
        write_report([], self.output_path)

        loaded_rows = read_csv(
            self.output_path,
            required_columns=[
                "record_id",
                "status",
                "column",
                "source_value",
                "target_value",
            ],
        )

        self.assertEqual(loaded_rows, [])


class TestReadXlsx(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.xlsx_path = Path(self.temp_dir.name) / "input.xlsx"

    def test_reads_text_numbers_and_blank_cells(self):
        workbook = Workbook()
        sheet = workbook.active

        sheet.append(["id", "customer", "amount"])
        sheet.append(["001", "Acme Supply", 150])
        sheet.append(["002", None, 275.5])

        workbook.save(self.xlsx_path)
        workbook.close()

        result = read_xlsx(
            self.xlsx_path,
            required_columns=["id", "customer", "amount"],
        )

        self.assertEqual(result, [
            {
                "id": "001",
                "customer": "Acme Supply",
                "amount": "150",
            },
            {
                "id": "002",
                "customer": "",
                "amount": "275.5",
            },
        ])

    def test_rejects_missing_required_column(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "customer"])

        workbook.save(self.xlsx_path)
        workbook.close()

        with self.assertRaisesRegex(
            ValueError, "missing required columns: amount"
        ):
            read_xlsx(self.xlsx_path, ["id", "amount"])

    def test_rejects_formula_cells(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "amount"])
        sheet.append(["A", "=10+5"])

        workbook.save(self.xlsx_path)
        workbook.close()

        with self.assertRaisesRegex(
            ValueError,
            "worksheet row 2 contains a formula or Excel error",
        ):
            read_xlsx(self.xlsx_path, ["id", "amount"])

    def test_skips_completely_blank_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "amount"])
        sheet.append(["A", 10])
        sheet.append([None, None])
        sheet.append(["B", 20])

        workbook.save(self.xlsx_path)
        workbook.close()

        result = read_xlsx(self.xlsx_path, ["id", "amount"])

        self.assertEqual(result, [
            {"id": "A", "amount": "10"},
            {"id": "B", "amount": "20"},
        ])

    def test_reads_selected_worksheet(self):
        workbook = Workbook()

        first_sheet = workbook.active
        first_sheet.title = "Overview"
        first_sheet.append(["note"])
        first_sheet.append(["Not the invoice table"])

        invoice_sheet = workbook.create_sheet("Invoices")
        invoice_sheet.append(["id", "amount"])
        invoice_sheet.append(["A", 150])

        workbook.save(self.xlsx_path)
        workbook.close()

        result = read_xlsx(
            self.xlsx_path,
            required_columns=["id", "amount"],
            sheet_name="Invoices",
        )

        self.assertEqual(result, [
            {"id": "A", "amount": "150"},
        ])

    def test_rejects_unknown_worksheet(self):
        workbook = Workbook()
        workbook.active.title = "Invoices"
        workbook.save(self.xlsx_path)
        workbook.close()

        with self.assertRaisesRegex(
            ValueError,
            "worksheet 'Missing' not found. Available worksheets: Invoices",
        ):
            read_xlsx(self.xlsx_path, sheet_name="Missing")


class TestReadTable(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.folder = Path(self.temp_dir.name)

    def test_reconciles_csv_against_xlsx(self):
        source_path = self.folder / "source.csv"
        target_path = self.folder / "target.XLSX"

        source_path.write_text(
            "id,amount\nA,150.00\nB,20.00\n",
            encoding="utf-8",
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "amount"])
        sheet.append(["B", 25])
        sheet.append(["A", 150])
        workbook.save(target_path)
        workbook.close()

        source_rows = read_table(source_path, ["id", "amount"])
        target_rows = read_table(target_path, ["id", "amount"])

        results = reconcile(
            source_rows,
            target_rows,
            key_column="id",
            comparison_columns=["amount"],
            numeric_columns=["amount"],
        )

        self.assertEqual(results, [
            {
                "record_id": "A",
                "status": "matched",
                "column": "",
                "source_value": "",
                "target_value": "",
            },
            {
                "record_id": "B",
                "status": "changed",
                "column": "amount",
                "source_value": "20.00",
                "target_value": "25",
            },
        ])

    def test_rejects_unsupported_extension(self):
        file_path = self.folder / "input.txt"

        with self.assertRaisesRegex(
            ValueError, "unsupported file extension"
        ):
            read_table(file_path)

    def test_passes_worksheet_selection_to_excel_reader(self):
        file_path = self.folder / "input.xlsx"

        workbook = Workbook()
        first_sheet = workbook.active
        first_sheet.title = "Overview"
        first_sheet.append(["note"])

        selected_sheet = workbook.create_sheet("Invoices")
        selected_sheet.append(["id", "amount"])
        selected_sheet.append(["A", 150])

        workbook.save(file_path)
        workbook.close()

        result = read_table(
            file_path,
            required_columns=["id", "amount"],
            sheet_name="Invoices",
        )

        self.assertEqual(result, [
            {"id": "A", "amount": "150"},
        ])

    def test_rejects_worksheet_selection_for_csv(self):
        file_path = self.folder / "input.csv"
        file_path.write_text(
            "id,amount\nA,150.00\n",
            encoding="utf-8",
        )

        with self.assertRaisesRegex(
            ValueError,
            "worksheet selection requires an XLSX file",
        ):
            read_table(
                file_path,
                required_columns=["id", "amount"],
                sheet_name="Invoices",
            )


class TestCommandLine(unittest.TestCase):
    def test_reconciles_csv_and_selected_excel_sheet(self):
        project_root = Path(__file__).resolve().parent.parent
        script_path = project_root / "src" / "reconcile.py"

        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            source_path = folder / "source.csv"
            target_path = folder / "target.xlsx"
            output_path = folder / "reports" / "result.csv"

            source_path.write_text(
                "id,amount\nA,150.00\nB,20.00\n",
                encoding="utf-8",
            )

            workbook = Workbook()
            workbook.active.title = "Overview"

            sheet = workbook.create_sheet("Invoice Data")
            sheet.append(["id", "amount"])
            sheet.append(["B", 25])
            sheet.append(["A", 150])

            workbook.save(target_path)
            workbook.close()

            completed = subprocess.run(
                [
                    sys.executable,
                    str(script_path),
                    "--source", str(source_path),
                    "--target", str(target_path),
                    "--target-sheet", "Invoice Data",
                    "--key", "id",
                    "--columns", "amount",
                    "--numeric-columns", "amount",
                    "--output", str(output_path),
                ],
                cwd=folder,
                capture_output=True,
                text=True,
                timeout=30,
            )

            self.assertEqual(
                completed.returncode,
                0,
                msg=completed.stderr,
            )
            self.assertIn("Matched records: 1", completed.stdout)
            self.assertIn("Changed records: 1", completed.stdout)
            self.assertIn("Source-only records: 0", completed.stdout)
            self.assertIn("Target-only records: 0", completed.stdout)

            self.assertEqual(read_csv(output_path), [
                {
                    "record_id": "A",
                    "status": "matched",
                    "column": "",
                    "source_value": "",
                    "target_value": "",
                },
                {
                    "record_id": "B",
                    "status": "changed",
                    "column": "amount",
                    "source_value": "20.00",
                    "target_value": "25",
                },
            ])