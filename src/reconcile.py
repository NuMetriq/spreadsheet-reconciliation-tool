import csv
import argparse
from pathlib import Path
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook

def read_csv(
    file_path: Path,
    required_columns: list[str] | None = None,
) -> list[dict[str, str]]:
    with file_path.open(newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        headers = reader.fieldnames

        if not headers:
            raise ValueError(f"{file_path}: missing header row")

        if len(headers) != len(set(headers)):
            raise ValueError(f"{file_path}: duplicate column names")

        if required_columns is not None:
            missing_columns = [
                column
                for column in required_columns
                if column not in headers
            ]

            if missing_columns:
                missing_names = ", ".join(missing_columns)
                raise ValueError(
                    f"{file_path}: missing required columns: {missing_names}"
                )

        rows = []

        for row_number, row in enumerate(reader, start=1):
            if None in row:
                raise ValueError(
                    f"{file_path}: data row {row_number} has extra fields"
                )

            if any(value is None for value in row.values()):
                raise ValueError(
                    f"{file_path}: data row {row_number} has missing fields"
                )

            rows.append(row)

        return rows


def read_xlsx(
    file_path: Path,
    required_columns: list[str] | None = None,
) -> list[dict[str, str]]:
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=False,
    )

    try:
        sheet = workbook.worksheets[0]
        worksheet_rows = sheet.iter_rows()
        header_cells = next(worksheet_rows, ())
        headers = [cell.value for cell in header_cells]

        if not headers:
            raise ValueError(f"{file_path}: missing header row")

        if any(
            not isinstance(header, str) or not header.strip()
            for header in headers
        ):
            raise ValueError(
                f"{file_path}: column names must be nonblank text"
            )

        if any(cell.data_type in ("f", "e") for cell in header_cells):
            raise ValueError(
                f"{file_path}: headers cannot contain formulas or errors"
            )

        if len(headers) != len(set(headers)):
            raise ValueError(f"{file_path}: duplicate column names")

        missing_columns = [
            column
            for column in (required_columns or [])
            if column not in headers
        ]

        if missing_columns:
            names = ", ".join(missing_columns)
            raise ValueError(
                f"{file_path}: missing required columns: {names}"
            )

        rows = []

        for row_number, cells in enumerate(worksheet_rows, start=2):
            if all(cell.value is None for cell in cells):
                continue

            values = []

            for cell in cells:
                if cell.data_type in ("f", "e"):
                    raise ValueError(
                        f"{file_path}: worksheet row {row_number} "
                        "contains a formula or Excel error"
                    )

                value = "" if cell.value is None else str(cell.value)
                values.append(value)

            if len(values) != len(headers):
                raise ValueError(
                    f"{file_path}: worksheet row {row_number} "
                    "does not match the header width"
                )

            rows.append(dict(zip(headers, values)))

        return rows

    finally:
        workbook.close()


def read_table(
    file_path: Path,
    required_columns: list[str] | None = None,
) -> list[dict[str, str]]:
    extension = file_path.suffix.lower()

    if extension == ".csv":
        return read_csv(file_path, required_columns)

    if extension == ".xlsx":
        return read_xlsx(file_path, required_columns)

    raise ValueError(
        f"{file_path}: unsupported file extension '{extension}'. "
        "Use .csv or .xlsx."
    )


def parse_decimal(value: str) -> Decimal:
    try:
        number = Decimal(value)
    except InvalidOperation:
        raise ValueError(f"Invalid numeric value: {value!r}") from None

    if not number.is_finite():
        raise ValueError(f"Numeric value must be finite: {value!r}")

    return number


def compare_rows(
    source_row: dict[str, str],
    target_row: dict[str, str],
    columns: list[str],
    numeric_columns: list[str] | None = None,
) -> dict[str, tuple[str, str]]:
    numeric_column_set = set(numeric_columns or [])
    differences = {}

    for column in columns:
        source_value = source_row[column]
        target_value = target_row[column]

        if column in numeric_column_set:
            values_match = (
                parse_decimal(source_value)
                == parse_decimal(target_value)
            )
        else:
            values_match = source_value == target_value

        if not values_match:
            differences[column] = (source_value, target_value)

    return differences

def index_rows(
    rows: list[dict[str, str]],
    key_column: str,
) -> dict[str, dict[str, str]]:
    indexed_rows = {}

    for row_number, row in enumerate(rows, start=1):
        if key_column not in row:
            raise ValueError(
                f"Data row {row_number}: missing key column '{key_column}'"
            )

        key = row[key_column]

        if key is None or not key.strip():
            raise ValueError(
                f"Data row {row_number}: blank key in '{key_column}'"
            )

        if key in indexed_rows:
            raise ValueError(f"Duplicate key found: {key}")

        indexed_rows[key] = row

    return indexed_rows


def reconcile(
    source_rows: list[dict[str, str]],
    target_rows: list[dict[str, str]],
    key_column: str,
    comparison_columns: list[str],
    numeric_columns: list[str] | None = None,
) -> list[dict[str, str]]:

    numeric_columns = numeric_columns or []

    unknown_columns = set(numeric_columns) - set(comparison_columns)

    if unknown_columns:
        names = ", ".join(sorted(unknown_columns))
        raise ValueError(
            f"Numeric columns must also be comparison columns: {names}"
        )
    
    source_index = index_rows(source_rows, key_column)
    target_index = index_rows(target_rows, key_column)

    all_ids = set(source_index) | set(target_index)
    results = []

    for record_id in sorted(all_ids):
        if record_id not in target_index:
            results.append({
                "record_id": record_id,
                "status": "source_only",
                "column": "",
                "source_value": "",
                "target_value": "",
            })
        elif record_id not in source_index:
            results.append({
                "record_id": record_id,
                "status": "target_only",
                "column": "",
                "source_value": "",
                "target_value": "",
            })
        else:
            differences = compare_rows(
                source_index[record_id],
                target_index[record_id],
                comparison_columns,
                numeric_columns=numeric_columns,
            )

            if differences:
                for column, values in differences.items():
                    source_value, target_value = values

                    results.append({
                        "record_id": record_id,
                        "status": "changed",
                        "column": column,
                        "source_value": source_value,
                        "target_value": target_value,
                    })
            else:
                results.append({
                    "record_id": record_id,
                    "status": "matched",
                    "column": "",
                    "source_value": "",
                    "target_value": "",
                })

    return results


def write_report(
    results: list[dict[str, str]],
    output_path: Path,
) -> None:
    fieldnames = [
        "record_id",
        "status",
        "column",
        "source_value",
        "target_value",
    ]

    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(results)


def summarize_results(
    results: list[dict[str, str]],
) -> dict[str, int]:
    record_ids_by_status: dict[str, set[str]] = {
        "matched": set(),
        "changed": set(),
        "source_only": set(),
        "target_only": set(),
    }

    for result in results:
        status = result["status"]
        record_id = result["record_id"]
        record_ids_by_status[status].add(record_id)

    return {
        status: len(record_ids)
        for status, record_ids in record_ids_by_status.items()
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description=(
            "Compare CSV or XLSX files and export a CSV reconciliation report."
        )
    )

    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--target", type=Path, required=True)
    parser.add_argument("--key", required=True)
    parser.add_argument("--columns", nargs="+", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--numeric-columns",
        nargs="+",
        default=[],
        help="Columns to compare as decimal numbers; must also appear in --columns.",
    )

    args = parser.parse_args()

    source_path = args.source.resolve()
    target_path = args.target.resolve()
    output_path = args.output.resolve()

    if output_path in (source_path, target_path):
        parser.error("Output path must differ from both input paths.")

    required_columns = [args.key] + args.columns

    try:
        source_rows = read_table(args.source, required_columns)
        target_rows = read_table(args.target, required_columns)

        results = reconcile(
            source_rows,
            target_rows,
            key_column=args.key,
            comparison_columns=args.columns,
            numeric_columns=args.numeric_columns,
        )

        write_report(results, args.output)

    except (ValueError, OSError, csv.Error) as error:
        parser.exit(status=1, message=f"Error: {error}\n")

    summary = summarize_results(results)

    print(f"Matched records: {summary['matched']}")
    print(f"Changed records: {summary['changed']}")
    print(f"Source-only records: {summary['source_only']}")
    print(f"Target-only records: {summary['target_only']}")
    print(f"\nReport saved to: {args.output}")